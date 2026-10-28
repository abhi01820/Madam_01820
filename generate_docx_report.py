"""Generate professional, beautifully formatted Word (.docx) document for HealthRisk AI."""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = Path(__file__).resolve().parent
DOCX_PATH = PROJECT_ROOT / "HealthRisk_AI_Project_Report.docx"


def set_cell_background(cell, hex_color: str):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding inside a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def create_project_report_docx():
    doc = Document()

    # Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)      # Deep Navy
    COLOR_ACCENT = RGBColor(0, 150, 136)       # Teal
    COLOR_SECONDARY = RGBColor(71, 85, 105)   # Slate gray
    COLOR_DARK = RGBColor(30, 41, 59)

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("HealthRisk AI")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Cardiometabolic Risk Analytics & Regional Nutrition Decision-Support System")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = COLOR_ACCENT

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(18)
    run_meta = meta_p.add_run("Technical Architecture • Model Evaluation • Clinical Decision Support • System Report")
    run_meta.font.name = "Calibri"
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = COLOR_SECONDARY

    # Divider Box: Educational Notice
    notice_tbl = doc.add_table(rows=1, cols=1)
    notice_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    notice_cell = notice_tbl.cell(0, 0)
    set_cell_background(notice_cell, "F1F5F9")
    set_cell_margins(notice_cell, top=140, bottom=140, left=200, right=200)
    np_par = notice_cell.paragraphs[0]
    np_par.paragraph_format.space_after = Pt(0)
    n_run1 = np_par.add_run("IMPORTANT CLINICAL NOTICE: ")
    n_run1.bold = True
    n_run1.font.size = Pt(9.5)
    n_run1.font.color.rgb = RGBColor(185, 28, 28)
    n_run2 = np_par.add_run(
        "HealthRisk AI is strictly an educational prototype and demonstration of applied machine learning pipelines. "
        "It is NOT a medical diagnostic tool or clinical medical device. All predictive risk categories represent synthetic "
        "model estimations, never absolute disease probabilities. Always consult a licensed physician for clinical interpretation."
    )
    n_run2.font.size = Pt(9.5)
    n_run2.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper function for Section Headings
    def add_section_header(text: str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY

    # Helper function for Subsections
    def add_subsection_header(text: str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_ACCENT

    # 1. Executive Summary
    add_section_header("1. Executive Summary & Problem Scope")
    p1 = doc.add_paragraph(
        "Cardiovascular diseases, hypertension, and type 2 diabetes represent major public health burdens globally and in South Asia. "
        "Early screening of multi-factorial markers (such as blood pressure, body mass index, lipid panels, glycemic levels, and lifestyle habits) "
        "is critical for preventive health awareness. However, traditional analytics tools frequently present complex numbers without context, "
        "rely on opaque black-box algorithms, or fail to provide culturally accessible dietary alternatives."
    )
    p1.paragraph_format.space_after = Pt(6)

    p2 = doc.add_paragraph(
        "HealthRisk AI was engineered as a self-contained, offline-runnable software application that bridges physiological validation, "
        "zero-leakage machine learning inference, and evidence-based nutrition science. It runs entirely on open-source Python technologies "
        "(Streamlit, Scikit-learn, Pandas, NumPy, Plotly, Joblib) without external API dependencies or paid cloud databases."
    )
    p2.paragraph_format.space_after = Pt(10)

    # 2. System Architecture & Diagram
    add_section_header("2. High-Level Architecture & Component Flow")
    doc.add_paragraph(
        "The system separates user input processing into distinct deterministic validation, emergency detection, machine learning classification, "
        "and rule-based recommendation pipelines:"
    )

    # Architecture Box Diagram (Table representation)
    arch_tbl = doc.add_table(rows=5, cols=3)
    arch_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Inches(1.8), Inches(0.4), Inches(4.3)]

    steps = [
        ("Step 1: Input Validation", "-->", "Deterministic checks against biological ranges (src/validation.py). Catches typos and out-of-range inputs."),
        ("Step 2: Emergency Triage", "-->", "Critical alert detection for Hypertensive Crisis (>=180/120 mmHg), severe hypoglycemia (<54 mg/dL), or severe hyperglycemia (>=300 mg/dL)."),
        ("Step 3: ML Inference", "-->", "Fitted Scikit-learn ColumnTransformer pipeline computes class likelihoods (Lower, Moderate, or Higher synthetic risk category)."),
        ("Step 4: Clinical Rules", "-->", "Formulates qualitative lifestyle and dietary advice based on WHO, ACC/AHA, NCEP ATP III, and ADA 2024 guidelines."),
        ("Step 5: Meal Generation", "-->", "Assembles daily South Indian meal plans complying with diet types (Veg/Vegan/Egg/Non-Veg) and declared food allergies.")
    ]

    for row_idx, (s_title, arr, s_desc) in enumerate(steps):
        row = arch_tbl.rows[row_idx]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width, c1.width, c2.width = col_widths[0], col_widths[1], col_widths[2]

        set_cell_background(c0, "0F172A" if row_idx % 2 == 0 else "1E293B")
        set_cell_background(c1, "F8FAFC")
        set_cell_background(c2, "F1F5F9" if row_idx % 2 == 0 else "E2E8F0")

        set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c1, top=60, bottom=60, left=40, right=40)
        set_cell_margins(c2, top=60, bottom=60, left=80, right=80)

        p_c0 = c0.paragraphs[0]
        r_c0 = p_c0.add_run(s_title)
        r_c0.font.bold = True
        r_c0.font.size = Pt(9.5)
        r_c0.font.color.rgb = RGBColor(255, 255, 255)

        p_c1 = c1.paragraphs[0]
        p_c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c1 = p_c1.add_run(arr)
        r_c1.font.bold = True
        r_c1.font.color.rgb = COLOR_ACCENT

        p_c2 = c2.paragraphs[0]
        r_c2 = p_c2.add_run(s_desc)
        r_c2.font.size = Pt(9)
        r_c2.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 3. Machine Learning Pipeline & Zero-Leakage Architecture
    add_section_header("3. Machine Learning Training Pipeline")
    doc.add_paragraph(
        "A rigorous, reproducible machine learning pipeline was constructed in src/train_model.py. "
        "To ensure data hygiene and patient privacy, a 2,500-sample synthetic dataset was generated with a fixed seed (random_state=42). "
        "A stratified 80/20 train/test split was executed, guaranteeing that test samples were never exposed during preprocessing."
    )

    doc.add_paragraph(
        "• Preprocessing: Numeric features (23 columns) were processed with SimpleImputer(strategy='median') and StandardScaler. "
        "Categorical features (Sex, Smoking Status, Alcohol Frequency) were encoded using SimpleImputer(strategy='most_frequent') and OneHotEncoder(handle_unknown='ignore')."
    )
    doc.add_paragraph(
        "• Model Comparison: A Logistic Regression baseline was compared against a Random Forest Classifier. Both were evaluated on the held-out 20% test set (N = 500)."
    )

    # Table of Evaluation Metrics
    add_subsection_header("Actual Test Holdout Evaluation Metrics (N = 500)")
    metrics_tbl = doc.add_table(rows=3, cols=6)
    metrics_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Model Candidate", "Test Accuracy", "F1 Score (Macro)", "F1 Score (Weighted)", "Precision (Macro)", "Recall (Macro)"]
    for i, h_text in enumerate(headers):
        cell = metrics_tbl.cell(0, i)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    data_rows = [
        ("Logistic Regression Baseline (Selected)", "81.60%", "0.8131", "0.8155", "83.17%", "79.92%"),
        ("Random Forest Classifier", "77.20%", "0.7082", "0.7645", "85.25%", "66.37%")
    ]

    for r_idx, row_data in enumerate(data_rows):
        row = metrics_tbl.rows[r_idx + 1]
        bg = "E6FFFA" if r_idx == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            p = cell.paragraphs[0]
            if c_idx > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.size = Pt(8.5)
            if r_idx == 0 and c_idx in [1, 2]:
                run.bold = True
                run.font.color.rgb = COLOR_ACCENT

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 4. Clinical Recommendation Engine
    add_section_header("4. Clinical Recommendation Engine & South Indian Nutrition")
    doc.add_paragraph(
        "The recommendation engine (src/recommendation_engine.py) applies clinical consensus thresholds without relying on third-party LLMs or fabricated macros:"
    )
    doc.add_paragraph(
        "1. Blood Pressure & Sodium Restriction: Adheres to 2017 ACC/AHA guidelines. For elevated pressures, it recommends reducing dietary sodium (< 2,000 mg/day) "
        "and seasoning with lemon, cumin, curry leaves, and coriander. It explicitly avoids unilaterally recommending high-potassium salt substitutes to protect renal safety."
    )
    doc.add_paragraph(
        "2. Glycemic Management: Adheres to ADA 2024 Standards of Care. Evaluates Fasting, Postprandial, Random glucose, and HbA1c individually without substituting markers. "
        "Recommends low-glycemic-index whole grains (finger millet/ragi, foxtail millet, hand-pounded red rice) paired with proteinaceous legumes."
    )
    doc.add_paragraph(
        "3. Lipid Management: Follows NCEP ATP III guidelines. Distinguishes between optimal, borderline, and high fractions for total cholesterol, LDL, HDL, and triglycerides. "
        "Focuses on soluble fiber (methi seeds, legumes, oats) and heart-smart unsaturated cooking fats."
    )
    doc.add_paragraph(
        "4. South Indian Cultural Meal Planner: Respects four dietary regimes (Vegetarian, Vegan, Eggetarian, Non-Vegetarian) and strictly filters out declared food allergens "
        "(peanuts, tree nuts, gluten, dairy/lactose, soy, shellfish, eggs). Descriptions are purely qualitative (portion mindfulness and nutritional mechanism), avoiding fabricated calorie numbers."
    )

    # 5. Application Pages
    add_section_header("5. Overview of Application Pages")
    pages_data = [
        ("1. Home Page", "Striking gradient hero, call-to-action buttons, quick-load demo profile presets, and an overview of all evaluated biometric pillars."),
        ("2. Health Assessment", "Tabbed form with live auto-calculating BMI, biological range validation, timing context selectors for blood glucose, and dietary preference options."),
        ("3. Prediction Results", "Emergency alert banner for acute readings, Plotly risk gauge, 5-pillar radar profile chart, metric cards with semantic status badges, and factor breakdowns."),
        ("4. Diet Planner", "Customizable daily meal plan (Breakfast, Lunch, Dinner, Snack) with regenerate button, local Indian grocery budget guidance, and cooking swaps."),
        ("5. Model Information", "Full transparency model card: synthetic data generation rubric, EDA findings, interactive confusion matrix heatmap, and documented limitations."),
        ("6. About & Disclaimer", "Prominent medical disclaimer, peer-reviewed clinical citations (WHO 2004, ACC/AHA 2017, NCEP ATP III, ADA 2024, ICMR-NIN), and technical architecture.")
    ]

    for p_title, p_desc in pages_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r_t = p.add_run(f"• {p_title}: ")
        r_t.bold = True
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = COLOR_DARK
        r_d = p.add_run(p_desc)
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_SECONDARY

    # 6. Verification & Test Suite
    add_section_header("6. Automated Verification & Test Coverage")
    doc.add_paragraph(
        "The codebase was verified with a comprehensive automated test suite (pytest -v) covering 19 distinct unit tests across validation, clinical rules, and model predictions:"
    )
    doc.add_paragraph("• tests/test_validation.py (7 tests): Plausible range checks, missing optional handling, physiological consistency (systolic > diastolic), and emergency detection.")
    doc.add_paragraph("• tests/test_recommendations.py (8 tests): Threshold evaluation (BMI, BP, Lipids, Glucose), sodium restriction, vegan/vegetarian safety, and allergy exclusion.")
    doc.add_paragraph("• tests/test_prediction.py (4 tests): Model initialization, output shape, probability sum checks, and prediction with missing optional fields.")
    doc.add_paragraph("Result: All 19 tests passed successfully in 5.0 seconds.")

    # 7. Local Run Instructions
    add_section_header("7. Local Setup & Execution Commands")
    cmd_tbl = doc.add_table(rows=3, cols=2)
    cmd_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, (action, cmd) in enumerate([
        ("Activate Virtual Environment", ".\\.venv\\Scripts\\Activate.ps1"),
        ("Run Automated Tests", "pytest -v"),
        ("Launch Streamlit App", "streamlit run app.py")
    ]):
        row = cmd_tbl.rows[r_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.5), Inches(4.0)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "0F172A")
        set_cell_margins(c0, top=50, bottom=50, left=80, right=80)
        set_cell_margins(c1, top=50, bottom=50, left=80, right=80)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(action)
        r0.font.bold = True
        r0.font.size = Pt(9)
        r0.font.color.rgb = COLOR_DARK

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(cmd)
        r1.font.name = "Consolas"
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor(0, 210, 186)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Save document
    doc.save(DOCX_PATH)
    print(f"Generated Word document at: {DOCX_PATH}")
    return DOCX_PATH


if __name__ == "__main__":
    create_project_report_docx()
